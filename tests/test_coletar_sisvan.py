import csv
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import coletar_sisvan as collector
import coletar_sisvan_derivado as helper


class SisvanCollectorTests(unittest.TestCase):
    def workbook(self, indicator, counts, phase=None):
        book = Workbook()
        sheet = book.active
        if phase is not None:
            sheet.cell(4, 1, "Fase da Vida: " + phase.official_label)
        sheet.cell(6, 1, indicator.official_title)
        width = 5 + 2 * len(indicator.categories) + 1
        offset = width
        for column, label in enumerate(collector.BASE_COLUMNS, offset + 1):
            sheet.cell(7, column, label)
        for index, category in enumerate(indicator.categories):
            sheet.cell(7, offset + 6 + 2 * index, category)
        sheet.cell(7, offset + width, "Total")
        total = sum(counts)
        base = ["NORTE", 11, "RO", 110001, "MUNICIPIO TESTE"]
        for column, value in enumerate(base, 1):
            sheet.cell(8, column, value)
        for index, count in enumerate(counts):
            sheet.cell(8, 6 + 2 * index, count)
            sheet.cell(8, 7 + 2 * index, "%s%%" % (count / total * 100))
        sheet.cell(8, width, total)
        output = io.BytesIO()
        book.save(output)
        book.close()
        return output.getvalue()

    def test_parse_imc_por_idade(self):
        indicator = collector.INDICATORS["imc_por_idade"]
        rows, schema = collector.parse_export(
            self.workbook(indicator, [1, 1, 5, 1, 1, 1]), "RO", indicator,
        )
        stats = collector.validate(rows, schema)
        self.assertEqual(stats["total"], 1)
        self.assertEqual(rows[0]["Magreza acentuada - Quantidade"], 1)
        self.assertEqual(rows[0]["Total"], 10)

    def test_parse_altura_por_idade(self):
        indicator = collector.INDICATORS["altura_por_idade"]
        rows, schema = collector.parse_export(
            self.workbook(indicator, [1, 2, 7]), "RO", indicator,
        )
        stats = collector.validate(rows, schema)
        self.assertEqual(stats["total"], 1)
        self.assertEqual(
            rows[0]["Altura Muito Baixa para a Idade - Quantidade"], 1,
        )
        self.assertEqual(rows[0]["Altura Adequada para a Idade - %"], "70.0%")

    def test_parse_peso_por_idade(self):
        indicator = collector.INDICATORS["peso_por_idade"]
        rows, schema = collector.parse_export(
            self.workbook(indicator, [1, 2, 6, 1]), "RO", indicator,
        )
        stats = collector.validate(rows, schema)
        self.assertEqual(stats["total"], 1)
        self.assertEqual(
            rows[0]["Peso Muito Baixo para a Idade - Quantidade"], 1,
        )
        self.assertEqual(rows[0]["Peso Adequado ou Eutrófico - %"], "60.0%")

    def test_peso_por_idade_mantem_esquema_em_todas_faixas(self):
        expected = collector.INDICATORS["peso_por_idade"].categories
        for age_key in collector.AGE_RANGES:
            query = collector.make_query(
                "crianca", "peso_por_idade", age_key, 2025,
            )
            self.assertEqual(query.indicator.code, "1")
            self.assertEqual(query.indicator.categories, expected)

    def default_args(self, *options):
        return collector.build_parser().parse_args([
            "--config", str(ROOT / "configuracoes/sisvan/coletas.json"), *options,
        ])

    def test_configuracao_padrao_usa_altura_e_peso_menores_5(self):
        args = self.default_args()
        queries = collector.resolve_queries(args)
        self.assertEqual(len(queries), 2)
        self.assertEqual(
            {(item.indicator.key, item.age_range.key) for item in queries},
            {
                ("altura_por_idade", "0_a_menor_5_anos"),
                ("peso_por_idade", "0_a_menor_5_anos"),
            },
        )
        for query in queries:
            self.assertEqual(query.year, 2025)
            self.assertEqual(query.phase.key, "crianca")
            self.assertEqual((query.age_range.start, query.age_range.end), ("0", "5"))
        self.assertEqual(len(collector.selected_states(args.ufs)), 27)
        for flag in (args.arquivo_unico, args.somar_faixas, args.harmonizar, args.populacao_geral):
            self.assertFalse(flag)
        outputs = [collector.output_path(args, q, collector.selected_states(args.ufs), len(queries)) for q in queries]
        self.assertEqual(len(set(outputs)), 2)
        self.assertEqual({p.parent.name for p in outputs}, {"altura_por_idade", "peso_por_idade"})

    def test_indice_explicito_substitui_as_duas_consultas_padrao(self):
        for index in ("altura_por_idade", "peso_por_idade"):
            with self.subTest(index=index):
                queries = collector.resolve_queries(self.default_args("--indices", index))
                self.assertEqual(len(queries), 1)
                self.assertEqual(queries[0].indicator.key, index)
                self.assertEqual(queries[0].age_range.key, "0_a_menor_5_anos")

    def test_ano_faixas_e_ufs_explicitos_continuam_configuraveis(self):
        args = self.default_args("--year", "2024", "--faixas-etarias", "0-5,5-10", "--ufs", "RO,AC")
        queries = collector.resolve_queries(args)
        self.assertEqual(len(queries), 4)
        self.assertEqual({q.year for q in queries}, {2024})
        self.assertEqual({q.age_range.key for q in queries}, {"0_a_menor_5_anos", "5_a_menor_10_anos"})
        self.assertEqual({q.indicator.key for q in queries}, {"altura_por_idade", "peso_por_idade"})
        self.assertEqual(collector.selected_states(args.ufs), [("RO", "11"), ("AC", "12")])

    def test_simulacao_padrao_nao_inicia_rede_nem_grava_produtos(self):
        output = io.StringIO()
        arguments = ["coletar_sisvan.py", "--config", str(ROOT / "configuracoes/sisvan/coletas.json"), "--dry-run"]
        with patch.object(sys, "argv", arguments), redirect_stdout(output), \
                patch.object(helper, "new_session") as session, \
                patch.object(collector, "Manifest") as manifest, \
                patch.object(collector, "write_csv") as writer:
            collector.main()
        self.assertIn("CRIANÇA | ALTURA X IDADE | 0 a < 5 anos | 2025", output.getvalue())
        self.assertIn("CRIANÇA | PESO X IDADE | 0 a < 5 anos | 2025", output.getvalue())
        session.assert_not_called()
        manifest.assert_not_called()
        writer.assert_not_called()

    def test_ambas_entradas_padrao_preservam_categorias_sem_dai_dpi(self):
        for query, counts, width in zip(collector.resolve_queries(self.default_args()), ([1, 2, 7], [1, 2, 6, 1]), (12, 14)):
            with self.subTest(index=query.indicator.key), tempfile.TemporaryDirectory() as directory:
                rows, schema = collector.parse_export(self.workbook(query.indicator, counts), "RO", query.indicator)
                stats = collector.validate(rows, schema)
                path = Path(directory) / (query.indicator.key + ".csv")
                collector.write_csv(path, rows, schema.columns)
                meta = collector.metadata_path(path)
                collector.write_metadata(meta, query, path, stats, [("RO", "11")], schema)
                self.assertEqual(len(schema.columns), width)
                self.assertEqual(rows[0]["Total"], sum(counts))
                self.assertTrue(all(rows[0][category + " - Quantidade"] == count
                                    for category, count in zip(query.indicator.categories, counts)))
                self.assertFalse(any("déficit" in c.lower() or c.lower() in {"dai", "dpi"} for c in schema.columns))
                self.assertEqual(json.loads(meta.read_text())["indicadores_derivados"], [])

    def test_payload_recebe_indice_e_faixa(self):
        payload = helper.report_payload(2025, "11", "3", "5", "10")
        self.assertEqual(payload["nu_indice_cri"], "3")
        self.assertEqual(payload["nu_idade_inicio"], "5")
        self.assertEqual(payload["nu_idade_fim"], "10")

    def test_configuracao_aceita_varias_faixas_por_indice(self):
        data = {
            "ano": 2025,
            "consultas": [
                {
                    "indice": "imc_por_idade",
                    "faixas_etarias": ["0-5", "5-10"],
                },
                {
                    "indice": "altura_por_idade",
                    "faixas_etarias": ["0-5", "5-10"],
                },
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.json"
            path.write_text(json.dumps(data), encoding="utf-8")
            args = SimpleNamespace(
                config=path, year=None, indices=None, faixas_etarias=None,
            )
            queries = collector.resolve_queries(args)
        self.assertEqual(len(queries), 4)
        self.assertEqual(
            {(item.indicator.key, item.age_range.key) for item in queries},
            {
                ("imc_por_idade", "0_a_menor_5_anos"),
                ("imc_por_idade", "5_a_menor_10_anos"),
                ("altura_por_idade", "0_a_menor_5_anos"),
                ("altura_por_idade", "5_a_menor_10_anos"),
            },
        )

    def test_todas_as_faixas_para_os_dois_indices(self):
        data = {
            "ano": 2025,
            "consultas": [
                {"indice": "imc_por_idade", "faixas_etarias": ["0-5"]},
                {"indice": "altura_por_idade", "faixas_etarias": ["0-5"]},
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.json"
            path.write_text(json.dumps(data), encoding="utf-8")
            args = SimpleNamespace(
                config=path, year=None, indices=None, faixas_etarias=None,
                todas_faixas=True,
            )
            queries = collector.resolve_queries(args)
        self.assertEqual(len(queries), 18)
        self.assertEqual(
            {item.age_range.key for item in queries}, set(collector.AGE_RANGES),
        )

    def test_arquivo_unico_usa_formato_longo(self):
        products = []
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            for key, counts in (
                ("imc_por_idade", [1, 1, 5, 1, 1, 1]),
                ("altura_por_idade", [1, 2, 7]),
            ):
                indicator = collector.INDICATORS[key]
                rows, schema = collector.parse_export(
                    self.workbook(indicator, counts), "RO", indicator,
                )
                output = directory / (key + ".csv")
                collector.write_csv(output, rows, schema.columns)
                products.append({
                    "query": collector.Query(
                        collector.PHASES["crianca"], indicator,
                        collector.AGE_RANGES["0_a_menor_5_anos"], 2025,
                    ),
                    "output": output,
                    "schema": schema,
                    "stats": collector.validate(rows, schema),
                })
            combined = directory / "combinado.csv"
            written = collector.write_combined(
                combined, products, [("RO", "11")],
            )
            with combined.open(newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(written, 9)
            self.assertEqual(len(rows), 9)
            self.assertEqual(
                {row["Índice antropométrico"] for row in rows},
                {"IMC X IDADE", "ALTURA X IDADE"},
            )
            self.assertEqual(
                {row["Faixa etária"] for row in rows}, {"0 a < 5 anos"},
            )
            self.assertTrue(combined.with_suffix(".metadados.json").exists())

    def test_payload_por_fase_da_vida(self):
        adolescent = helper.report_payload(
            2025, "11", life_cycle="2", adolescent_index_code="2",
        )
        self.assertEqual(adolescent["nu_ciclo_vida"], "2")
        self.assertEqual(adolescent["nu_indice_ado"], "2")
        pregnant = helper.report_payload(2025, "11", life_cycle="5")
        self.assertEqual(pregnant["nu_ciclo_vida"], "5")
        self.assertEqual(pregnant["nu_idade_ges"], "99")

    def test_populacao_geral_usa_fases_exclusivas(self):
        queries = collector.population_queries(2025, include_pregnant=True)
        self.assertEqual(
            [item.phase.key for item in queries],
            ["crianca", "crianca", "adolescente", "adulto", "idoso", "gestante"],
        )
        children = [
            item.age_range.key for item in queries if item.phase.key == "crianca"
        ]
        self.assertEqual(
            children, ["0_a_menor_5_anos", "5_a_menor_10_anos"],
        )
        collector.validate_population_queries(queries)

    def test_soma_rejeita_faixas_sobrepostas(self):
        queries = [
            collector.make_query("crianca", "imc_por_idade", "0-5", 2025),
            collector.make_query("crianca", "imc_por_idade", "2-5", 2025),
        ]
        with self.assertRaisesRegex(RuntimeError, "faixas sobrepostas"):
            collector.validate_nonoverlapping_ranges(queries)

    def test_soma_aceita_faixas_contiguas(self):
        queries = [
            collector.make_query("crianca", "imc_por_idade", "0-5", 2025),
            collector.make_query("crianca", "imc_por_idade", "5-10", 2025),
        ]
        collector.validate_nonoverlapping_ranges(queries)

    def test_esquemas_oficiais_mudam_por_fase(self):
        older_child = collector.make_query(
            "crianca", "imc_por_idade", "5-10", 2025,
        )
        adolescent = collector.make_query(
            "adolescente", "imc_por_idade", None, 2025,
        )
        adult = collector.make_query("adulto", "imc", None, 2025)
        self.assertIn("Obesidade grave (5-10 anos)", older_child.indicator.categories)
        self.assertIn("Obesidade Grave", adolescent.indicator.categories)
        self.assertIn("Obesidade Grau III", adult.indicator.categories)

    def test_fase_da_resposta_e_validada(self):
        query = collector.make_query(
            "adolescente", "imc_por_idade", None, 2025,
        )
        binary = self.workbook(
            query.indicator, [1, 1, 5, 1, 1, 1], query.phase,
        )
        rows, _ = collector.parse_export(
            binary, "RO", query.indicator, query.phase,
        )
        self.assertEqual(len(rows), 1)
        with self.assertRaisesRegex(RuntimeError, "nao corresponde a fase"):
            collector.parse_export(
                binary, "RO", query.indicator, collector.PHASES["adulto"],
            )

    def test_soma_recalcula_percentual_e_registra_fontes(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            products = []
            for position, (age, counts) in enumerate((
                ("0-5", [1, 1, 4, 1, 1, 2]),
                ("5-10", [1, 1, 4, 1, 1, 2]),
            )):
                query = collector.make_query(
                    "crianca", "imc_por_idade", age, 2025,
                )
                rows, schema = collector.parse_export(
                    self.workbook(query.indicator, counts), "RO", query.indicator,
                )
                source = directory / ("source_%s.csv" % position)
                collector.write_csv(source, rows, schema.columns)
                products.append({
                    "query": query, "output": source, "schema": schema,
                    "stats": collector.validate(rows, schema),
                })
            args = SimpleNamespace(output_dir=directory / "output")
            output = collector.write_summed_product(
                args, products, [("RO", "11")],
            )
            with output.open(newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))
            severe = next(
                row for row in rows
                if row["Classificação nutricional oficial"] == "Magreza acentuada"
            )
            metadata = json.loads(
                output.with_suffix(".metadados.json").read_text(encoding="utf-8")
            )
        self.assertEqual(severe["Quantidade"], "2")
        self.assertEqual(severe["Total"], "20")
        self.assertEqual(severe["Percentual recalculado"], "10.0")
        self.assertTrue(metadata["sobreposicao_validada"])
        self.assertEqual(len(metadata["fontes"]), 2)

    def test_harmonizacao_preserva_componentes_e_denominadores(self):
        plans = [
            ("crianca", "imc_por_idade", "0-5", [1, 1, 4, 1, 1, 2]),
            ("crianca", "imc_por_idade", "5-10", [1, 1, 4, 1, 1, 2]),
            ("adolescente", "imc_por_idade", None, [1, 1, 5, 1, 1, 1]),
            ("adulto", "imc", None, [2, 4, 1, 1, 1, 1]),
            ("idoso", "imc", None, [2, 5, 3]),
        ]
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            products = []
            for position, (phase, index, age, counts) in enumerate(plans):
                query = collector.make_query(phase, index, age, 2025)
                rows, schema = collector.parse_export(
                    self.workbook(query.indicator, counts), "RO", query.indicator,
                )
                output = directory / ("source_%s.csv" % position)
                collector.write_csv(output, rows, schema.columns)
                products.append({
                    "query": query,
                    "output": output,
                    "schema": schema,
                    "stats": collector.validate(rows, schema),
                })
            mapping_data, mapping = collector.load_harmonization(
                ROOT / "configuracoes/sisvan/harmonizacao_v1.json",
            )
            self.assertEqual(mapping_data["versao"], "1.0.0")
            phase_results = []
            for phase in ("crianca", "adolescente", "adulto", "idoso"):
                group = [item for item in products if item["query"].phase.key == phase]
                phase_key, year, facts = collector.build_phase_facts(group, mapping)
                phase_results.append({
                    "phase": phase_key, "year": year, "facts": facts,
                    "output": group[0]["output"],
                })
            year, facts = collector.combine_general_facts(phase_results)
        fact = facts["110001"]
        self.assertEqual(year, 2025)
        self.assertEqual(fact["magreza_acentuada"], 3)
        self.assertEqual(fact["magreza"], 3)
        self.assertEqual(fact["baixo_peso"], 4)
        self.assertEqual(fact["deficit"], 10)
        self.assertEqual(fact["denom_crianca_adolescente"], 30)
        self.assertEqual(fact["denom_adulto_idoso"], 20)
        self.assertEqual(fact["denom_risco"], 10)
        self.assertEqual(fact["total"], 50)

    def test_preserva_celulas_decimais_sem_calcular_deficit(self):
        indicator = collector.INDICATORS["altura_por_idade"]
        binary = self.workbook(indicator, [26, 62, 1020])
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(binary))
        sheet = book.active
        sheet.cell(8, 10, 1.02)
        sheet.cell(8, 12, 1.108)
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        rows, schema = collector.parse_export(stream.getvalue(), "RO", indicator)
        stats = collector.validate(rows, schema)
        self.assertEqual(rows[0]["Altura Adequada para a Idade - Quantidade"], 1.02)
        self.assertEqual(rows[0]["Total"], 1.108)
        self.assertEqual(len(schema.columns), 12)
        self.assertFalse(any("Déficit" in c for c in schema.columns))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "altura.csv"
            collector.write_csv(path, rows, schema.columns)
            query = collector.make_query("crianca", "altura_por_idade", "0-5", 2025)
            meta = path.with_suffix(".metadados.json")
            collector.write_metadata(meta, query, path, stats, [("RO", "11")], schema)
            with path.open(encoding="utf-8-sig", newline="") as handle:
                row = next(csv.DictReader(handle))
            self.assertEqual(row["Total"], "1.108")
            self.assertEqual(row["Altura Adequada para a Idade - Quantidade"], "1.02")
            metadata = json.loads(meta.read_text())
            self.assertEqual(metadata["indicadores_derivados"], [])
            self.assertEqual(metadata["sha256_csv"], collector.sha256(path.read_bytes()))

    def test_main_nao_calcula_indicador_infantil_automaticamente(self):
        source = Path(collector.__file__).read_text()
        self.assertNotIn("write_under5_growth_product", source)
        self.assertNotIn("Déficit de estatura", source)
        self.assertNotIn("Déficit de peso para idade", source)

    def test_faixa_invalida_e_rejeitada(self):
        with self.assertRaises(ValueError):
            collector.canonical_age("3-9")


if __name__ == "__main__":
    unittest.main()
