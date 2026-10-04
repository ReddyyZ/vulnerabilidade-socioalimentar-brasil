import ast
import base64
import contextlib
import hashlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import analise_sobreposicao as analysis
from formatacao_tabelas import tabela_br


class OverlapTests(unittest.TestCase):
    def sample(self):
        return pd.DataFrame({
            "codigo_ibge_6": ["110001", "110002", "110003", "110004", "110005"],
            "ivs": [0.1, 0.2, 0.3, 0.45, np.nan],
            "idhm": [0.9, 0.8, 0.7, 0.55, np.nan],
            "cadinsan_pct_com_PBF": [10., 20., 30., 40., np.nan],
            "cadinsan_pct_sem_PBF": [20., 30., 40., 50., np.nan],
            "cadinsan_n_com_PBF": [10., 20., 30., 40., np.nan],
            "cadinsan_n_sem_PBF": [20., 30., 40., 50., np.nan],
            "cadastros_cadunico_cadinsan": [100., 100., 100., 100., np.nan],
            "dai_pct": [1., 2., 3., 4., np.nan],
            "avaliados_altura": [100., 100., 100., 100., 0.],
        })

    def test_missing_information_is_not_low_risk(self):
        result, _ = analysis.classify(self.sample())
        self.assertTrue(pd.isna(result.loc[4, "prioritario"]))
        self.assertEqual(result.loc[4, "perfil"], "Informação insuficiente")
        self.assertIn("sem avaliações de altura", result.loc[4, "motivo_nao_classificacao"])
        self.assertTrue(result.loc[3, "prioritario"])

    def test_ties_and_minimum_boundary_are_included(self):
        frame = self.sample()
        for column in ("ivs", "cadinsan_pct_com_PBF", "dai_pct"):
            frame.loc[2, column] = frame.loc[3, column]
        frame.loc[2, "idhm"] = frame.loc[3, "idhm"]
        result, _ = analysis.classify(frame)
        self.assertEqual(int(result.prioritario.fillna(False).sum()), 2)
        frame.loc[3, "avaliados_altura"] = 99
        result, _ = analysis.classify(frame)
        self.assertFalse(result.loc[3, "elegivel_principal"])
        self.assertTrue(pd.isna(result.loc[3, "criterio_dai"]))

    def test_low_idhm_is_required_and_high_idhm_blocks_selection(self):
        frame = self.sample()
        frame.loc[3, "idhm"] = 0.99
        result, _ = analysis.classify(frame)
        self.assertFalse(result.loc[3, "prioritario"])
        self.assertFalse(result.loc[3, "criterio_idhm"])
        self.assertTrue(result.loc[3, "elegivel_principal"])
        self.assertEqual(result.loc[3, "numero_criterios_primarios"], 3)
        self.assertEqual(result.loc[3, "perfil"], "DAI elevado sem convergência social-alimentar")

    def test_missing_idhm_is_insufficient_information(self):
        frame = self.sample()
        frame.loc[3, "idhm"] = np.nan
        result, _ = analysis.classify(frame)
        self.assertFalse(result.loc[3, "elegivel_principal"])
        self.assertTrue(pd.isna(result.loc[3, "prioritario"]))
        self.assertTrue(pd.isna(result.loc[3, "numero_criterios_primarios"]))
        self.assertEqual(result.loc[3, "perfil"], "Informação insuficiente")
        self.assertEqual(result.loc[3, "motivo_nao_classificacao"], "IDHM ausente")

    def test_fixed_social_cutoffs_include_ivs_and_exclude_idhm_boundary(self):
        frame = self.sample()
        frame.loc[:3, "ivs"] = [0.400, 0.401, 0.402, 0.45]
        frame.loc[:3, "idhm"] = [0.599, 0.600, 0.601, 0.55]
        result, cuts = analysis.classify(frame)
        indexed = cuts.set_index("indicador")
        self.assertEqual(indexed.loc["ivs", "corte"], 0.401)
        self.assertEqual(indexed.loc["ivs", "operador"], ">=")
        self.assertEqual(indexed.loc["idhm", "corte"], 0.600)
        self.assertEqual(indexed.loc["idhm", "operador"], "<")
        for indicator in ("ivs", "idhm"):
            self.assertEqual(indexed.loc[indicator, "metodo"], "fixo")
            self.assertIsNone(indexed.loc[indicator, "percentil"])
        self.assertEqual(result.loc[:3, "criterio_ivs"].tolist(), [False, True, True, True])
        self.assertEqual(result.loc[:3, "criterio_idhm"].tolist(), [True, False, False, True])
        self.assertEqual(result.loc[3, "numero_criterios_primarios"], 4)
        self.assertEqual(result.loc[3, "perfil"], "Convergência dos quatro critérios")

    def test_quantile_changes_do_not_change_social_criteria(self):
        frame = self.sample()
        reference, cuts75 = analysis.classify(frame, quantile=0.75)
        result, cuts80 = analysis.classify(frame, quantile=0.8)
        for indicator in ("ivs", "idhm"):
            pd.testing.assert_series_equal(reference[f"criterio_{indicator}"], result[f"criterio_{indicator}"])
            pd.testing.assert_series_equal(cuts75.set_index("indicador").loc[indicator],
                                           cuts80.set_index("indicador").loc[indicator])
        for indicator in ("cadinsan", "dai"):
            self.assertEqual(cuts75.set_index("indicador").loc[indicator, "metodo"], "quantil")
            self.assertLess(cuts75.set_index("indicador").loc[indicator, "corte"],
                            cuts80.set_index("indicador").loc[indicator, "corte"])
        json.dumps(cuts75.to_dict("records"), allow_nan=False)

    def test_numeric_formats_and_invalid_input(self):
        parsed = analysis.numeric(pd.Series(["18,7%", "2.35%", "", None]), "teste")
        self.assertAlmostEqual(parsed.iloc[0], 18.7)
        self.assertAlmostEqual(parsed.iloc[1], 2.35)
        self.assertTrue(parsed.iloc[2:].isna().all())
        with self.assertRaises(ValueError):
            analysis.numeric(pd.Series(["sem dados"]), "teste")

    def test_zero_cadinsan_denominator_is_ineligible(self):
        frame = self.sample()
        frame.loc[3, "cadastros_cadunico_cadinsan"] = 0
        result, _ = analysis.classify(frame)
        self.assertTrue(pd.isna(result.loc[3, "criterio_cadinsan"]))
        self.assertFalse(result.loc[3, "elegivel_principal"])

    def test_regional_percent_uses_ratio_of_sums(self):
        frame = pd.DataFrame({
            "regiao": ["Norte", "Norte"], "elegivel_principal": [True, True],
            "prioritario": [True, False], "dai_n": [1, 90], "avaliados_altura": [10, 100],
        })
        result = analysis.regional_summary(frame).iloc[0]
        self.assertAlmostEqual(result.dai_pct_agregado, 91 / 110 * 100)
        self.assertNotAlmostEqual(result.dai_pct_agregado, 50.)

    def test_sensitivity_contains_reference_and_keeps_missing(self):
        result, _ = analysis.classify(self.sample())
        table, stability = analysis.sensitivity(self.sample(), result)
        principal = table.query("cenario == 'com_PBF' and quantil == 0.75 and minimo_avaliados == 100").iloc[0]
        self.assertEqual(principal.jaccard_principal, 1.)
        self.assertEqual(stability.set_index("codigo_ibge_6").loc["110005", "cenarios_elegiveis"], 0)
        self.assertEqual(len(table), 12)
        self.assertTrue(table.corte_ivs.eq(0.401).all())
        self.assertTrue(table.corte_idhm.eq(0.600).all())
        self.assertTrue(table.operador_idhm.eq("<").all())

    def test_invalid_parameters_fail(self):
        for kwargs in ({"quantile": 1}, {"min_evaluated": 0}, {"scenario": "inventado"}):
            with self.assertRaises(ValueError):
                analysis.classify(self.sample(), **kwargs)


class ResearchSnapshotTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base, cls.audit, cls.hashes = analysis.prepare_base(ROOT / "dados/pesquisa")

    def test_preserves_municipal_union_and_source_presence(self):
        self.assertEqual(len(self.base), 5571)
        self.assertFalse(self.base.codigo_ibge_6.duplicated().any())
        self.assertEqual(self.base.tem_ivs_idhm.sum(), 5565)
        self.assertEqual(self.base.tem_cadunico.sum(), 5564)
        self.assertEqual(self.base.tem_cadinsan.sum(), 5570)
        self.assertEqual(len(self.hashes), 4)

    def test_deficits_match_source_counts_and_zero_is_missing(self):
        self.assertEqual(self.base.dai_n.sum(), 877708)
        zero = self.base.loc[self.base.codigo_ibge_6 == "510183"].iloc[0]
        self.assertTrue(pd.isna(zero.dai_pct))
        self.assertEqual(zero.avaliados_altura, 0.)
        self.assertTrue(pd.isna(zero.codigo_ibge_7))

    def test_default_selection_and_thresholds(self):
        result, cuts = analysis.classify(self.base)
        self.assertEqual(result.elegivel_principal.sum(), 5326)
        self.assertEqual(result.prioritario.fillna(False).sum(), 242)
        self.assertAlmostEqual(cuts.set_index("indicador").loc["ivs", "corte"], 0.401)
        self.assertAlmostEqual(cuts.set_index("indicador").loc["cadinsan", "corte"], 10.969776400916132)
        self.assertAlmostEqual(cuts.set_index("indicador").loc["dai", "corte"], 12.380572037121558)
        self.assertAlmostEqual(cuts.set_index("indicador").loc["idhm", "corte"], 0.600)
        selected = result.loc[result.prioritario.fillna(False)]
        self.assertTrue(selected.criterio_idhm.all())
        self.assertTrue(selected.numero_criterios_primarios.eq(4).all())

    def test_four_criteria_require_low_idhm_with_fixed_ivs(self):
        result, _ = analysis.classify(self.base)
        three = result[["criterio_ivs", "criterio_cadinsan", "criterio_dai"]].fillna(False).all(axis=1)
        four = result.prioritario.fillna(False)
        self.assertEqual(three.sum(), 342)
        self.assertFalse((four & ~three).any())
        self.assertEqual((three & ~four).sum(), 100)
        self.assertFalse(result.loc[three & ~four, "criterio_idhm"].any())

    def test_cadinsan_uses_unrounded_ratios_and_preserves_source(self):
        indexed = self.base.set_index("codigo_ibge_6")
        city = indexed.loc["330455"]
        self.assertEqual(city.cadinsan_pct_com_PBF_arquivo, 18.7)
        self.assertAlmostEqual(city.cadinsan_pct_com_PBF, 100 * 97708 / 521373)
        self.assertNotEqual(city.cadinsan_pct_com_PBF, city.cadinsan_pct_com_PBF_arquivo)
        self.assertAlmostEqual(city.cadinsan_diferenca_pp_com_PBF,
                               city.cadinsan_pct_com_PBF - 18.7)
        for scenario in ("com_PBF", "sem_PBF"):
            expected = 100 * self.base[f"cadinsan_n_{scenario}"] / self.base.cadastros_cadunico_cadinsan
            np.testing.assert_allclose(self.base[f"cadinsan_pct_{scenario}"], expected, equal_nan=True)
        zero = indexed.loc["510183"]
        self.assertTrue(pd.isna(zero.cadinsan_pct_com_PBF))

    def test_rounding_comparison_explains_selection_change(self):
        result, _ = analysis.classify(self.base)
        comparison, old_cuts = analysis.compare_rounding(self.base, result)
        self.assertEqual(comparison.selecionado_percentual_csv.fillna(False).sum(), 243)
        self.assertEqual(comparison.selecionado_sem_arredondamento.fillna(False).sum(), 242)
        changed = comparison.loc[comparison.mudou_selecao]
        self.assertEqual(changed.codigo_ibge_6.tolist(), ["292410"])
        self.assertAlmostEqual(old_cuts.set_index("indicador").loc["cadinsan", "corte"], 11.)

    def test_catalogue_matches_social_sources_and_defines_units(self):
        table, meta = analysis.validate_catalogue(ROOT / "dataset_catalogue.json", ROOT / "dados/pesquisa")
        self.assertEqual(len(table), 3)
        self.assertTrue(table.corresponde_ao_catalogo.all())
        self.assertEqual(meta["cadunico"]["unidade"], "pessoas")
        self.assertEqual(meta["cadunico"]["periodo"], "2026-06")
        self.assertEqual(meta["cadinsan"]["unidade"], "famílias")
        self.assertEqual(meta["cadinsan"]["referencia_relatorio_oficial"], "2025-01")
        pd.testing.assert_series_equal(self.base.cadunico_pessoas_2026_06,
                                       self.base.cadunico_valor_original, check_names=False)

    def test_catalogue_rejects_hash_mismatch(self):
        data = json.loads((ROOT / "dataset_catalogue.json").read_text())
        data["datasets"]["municipios_cadunico"]["stats"]["checkSum"] = "sha256:" + "0" * 64
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "catalogue.json"
            path.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Catálogo não corresponde"):
                analysis.validate_catalogue(path, ROOT / "dados/pesquisa")

    def test_notebook_is_valid_and_payload_preserves_inputs(self):
        import nbformat
        notebook = nbformat.read(ROOT / "notebooks/01_sobreposicao_criterios.ipynb", as_version=4)
        nbformat.validate(notebook)
        code = [cell.source for cell in notebook.cells if cell.cell_type == "code"]
        for source in code:
            ast.parse(source)
        source = next(c for c in code if "PACOTE_BASE64 =" in c)
        tree = ast.parse(source)
        encoded = next(n.value.value for n in tree.body if isinstance(n, ast.Assign)
                       and isinstance(n.targets[0], ast.Name) and n.targets[0].id == "PACOTE_BASE64")
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(encoded))) as archive:
            for relative in analysis.DATA_FILES.values():
                self.assertEqual(archive.read(relative), (ROOT / "dados/pesquisa" / relative).read_bytes())
            self.assertEqual(archive.read("documentacao/dataset_catalogue.json"),
                             (ROOT / "dataset_catalogue.json").read_bytes())
            meta = json.loads(archive.read("apoio/ibge/malha_municipal_simplificada.metadados.json"))
            mesh_bytes = archive.read("apoio/ibge/malha_municipal_simplificada.geojson")
            self.assertEqual(hashlib.sha256(mesh_bytes).hexdigest(), meta["sha256"])
            audit = analysis.geometry_audit(json.loads(mesh_bytes), self.base)
            self.assertTrue(audit.tem_geometria.all())

    def test_notebook_explains_sources_and_method_without_project_documents(self):
        import nbformat
        notebook = nbformat.read(ROOT / "notebooks/01_sobreposicao_criterios.ipynb", as_version=4)
        first_code = next(i for i, cell in enumerate(notebook.cells) if cell.cell_type == "code")
        opening = "\n".join(cell.source for cell in notebook.cells[:first_code])
        markdown = "\n".join(cell.source for cell in notebook.cells if cell.cell_type == "markdown")
        sources = json.loads((ROOT / "fonte_dados.json").read_text())
        self.assertIn("Bases de dados e procedência", opening)
        self.assertIn("Pergunta da pesquisa", opening)
        for filename in ("atlasivs_municipios_2010.csv", "municipios-cadunico.json",
                         "CADINSAN_2025_dados_municipais.csv", "sisvan_relatorios"):
            self.assertIn(sources[filename], opening)
        for text in ("2010", "Junho/2026", "Janeiro/2025", "0 a menos de 5 anos",
                     "pessoas cadastradas", "famílias", "API v4 do IBGE", "não extraídas diretamente"):
            self.assertIn(text, opening)
        for text in ("3391236", "scripts/", "docs/metodologia/", "dados/pesquisa/", "No projeto:"):
            self.assertNotIn(text, markdown)
        self.assertNotIn("DPI", markdown)
        for text in ("DAI", "0,401", "0,600", "100 avaliações", "Jaccard", "causalidade"):
            self.assertIn(text, markdown)
        parameter_cell = notebook.cells[first_code]
        self.assertIn("QUANTIL = 0.75", parameter_cell.source)
        self.assertIn("MINIMO_AVALIADOS = 100", parameter_cell.source)
        self.assertNotIn("CODIGO_ANALISE_SHA256", parameter_cell.source)
        setup = next(cell for cell in notebook.cells if "CODIGO_ANALISE_SHA256 =" in cell.source)
        self.assertEqual(setup.cell_type, "code")
        self.assertTrue(setup.metadata["jupyter"]["source_hidden"])
        self.assertEqual(sum(cell.cell_type == "code" for cell in notebook.cells), 13)

    def test_section_eight_separates_selection_from_context_without_changing_data(self):
        import nbformat
        notebook = nbformat.read(ROOT / "notebooks/01_sobreposicao_criterios.ipynb", as_version=4)
        section = next(cell for cell in notebook.cells
                       if cell.cell_type == "code" and "indicadores_selecao =" in cell.source)
        classified, _ = analysis.classify(self.base)
        original = classified.copy(deep=True)
        shown = []
        namespace = {"classificados": classified, "display": shown.append,
                     "Markdown": lambda text: text, "regional_summary": analysis.regional_summary,
                     "tabela_br": tabela_br,
                     "CENARIO_CADINSAN": "com_PBF"}
        with contextlib.redirect_stdout(io.StringIO()):
            exec(section.source, namespace)
        from pandas.io.formats.style import Styler
        tables = [value.data for value in shown if isinstance(value, Styler)]
        self.assertEqual(len(tables), 3)
        selection, context, regions = tables
        self.assertEqual(selection.columns.tolist(), [
            "Município", "UF", "IVS", "IDHM", "CadInsan (%)", "DAI (%)",
            "Famílias no universo CadInsan", "Avaliações de altura"])
        self.assertEqual(context.columns.tolist(), [
            "Município", "UF", "Pessoas no CadÚnico — jun/2026",
            "Famílias em risco estimado — jan/2025"])
        self.assertEqual(len(selection), 30)
        pd.testing.assert_frame_equal(selection[["Município", "UF"]], context[["Município", "UF"]])
        pd.testing.assert_frame_equal(classified, original)
        pd.testing.assert_frame_equal(regions, analysis.regional_summary(classified))
        priority = namespace["prioritarios"]
        self.assertEqual(len(priority), 242)
        for column in ("criterio_idhm",
                       "codigo_ibge_6", "codigo_ibge_7"):
            self.assertIn(column, priority.columns)
            self.assertNotIn(column, namespace["colunas_selecao"])
            self.assertNotIn(column, namespace["colunas_caracterizacao"])
        pd.testing.assert_series_equal(selection["DAI (%)"], priority.dai_pct.head(30), check_names=False)
        pd.testing.assert_series_equal(context["Pessoas no CadÚnico — jun/2026"],
                                       priority.cadunico_pessoas_2026_06.head(30), check_names=False)
        html = next(value.to_html() for value in shown if isinstance(value, Styler)
                    and "Pessoas no CadÚnico — jun/2026" in value.data.columns)
        self.assertIn(">9.486</td>", html)
        self.assertIn(">223</td>", html)
        self.assertNotIn("9486.000", html)

    def test_height_only_input_and_analysis_do_not_export_weight(self):
        sources, _ = analysis.read_sources(ROOT / "dados/pesquisa")
        self.assertEqual(len(sources["sisvan"].columns), 12)
        for column in sources["sisvan"]:
            self.assertNotIn("Déficit", column)
            self.assertNotIn("Peso", column)
        self.assertNotIn("Ano", sources["sisvan"])
        result, cuts = analysis.classify(self.base)
        self.assertEqual(set(cuts.indicador), {"ivs", "idhm", "cadinsan", "dai"})
        for column in result:
            self.assertFalse(column.startswith(("dpi", "peso")))
        city = self.base.set_index("codigo_ibge_6").loc["110001"]
        self.assertEqual(city.altura_adequada_n_valor_fonte, "1.02")
        self.assertEqual(city.avaliados_altura_valor_fonte, "1.108")
        self.assertEqual(city.altura_adequada_n, 1020)
        self.assertEqual(city.avaliados_altura, 1108)
        self.assertAlmostEqual(city.dai_pct, 100 * (26 + 62) / 1108)

    def test_new_csv_preserves_every_source_cell_and_raw_hash(self):
        from openpyxl import load_workbook
        metadata = analysis.sisvan_metadata(ROOT / "dados/pesquisa")
        source_path = ROOT / "dados/pesquisa" / analysis.DATA_FILES["sisvan"]
        import csv
        with source_path.open(encoding="utf-8-sig", newline="") as handle:
            rows = {row["Código IBGE"]: row for row in csv.DictReader(handle)}
        self.assertEqual(len(metadata["xlsx_originais"]), 27)
        observed = set()
        for item in metadata["xlsx_originais"]:
            path = ROOT / item["arquivo_local"]
            content = path.read_bytes()
            self.assertEqual(hashlib.sha256(content).hexdigest(), item["hash_arquivo"])
            filters = json.loads(item["filtros"])
            self.assertEqual(filters["nu_indice_cri"], "3")
            self.assertEqual(filters["nuAno"], "2025")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                for values in book.active.iter_rows(values_only=True):
                    if len(values) < 12 or str(values[3]) not in rows:
                        continue
                    code = str(values[3])
                    observed.add(code)
                    for column, value in zip(metadata["colunas"], values[:12]):
                        self.assertEqual(rows[code][column], str(value), (code, column))
            finally:
                book.close()
        self.assertEqual(observed, set(rows))
        csv_origin = ROOT / metadata["arquivo"]
        self.assertEqual(source_path.read_bytes(), csv_origin.read_bytes())

    def test_historical_combined_input_is_preserved(self):
        path = ROOT / "dados/historico/pesquisa_sisvan_altura_peso_2025/indicadores_altura_peso_idade_menores_5_2025.csv"
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),
                         "21aa902b60fd5f08b1109f95a1c7e9a635a937fc96661fb57fe2519dcca9a92d")

    def test_count_interpretation_preserves_coherent_integers_and_rejects_inconsistency(self):
        values, changed, solutions = analysis.interpret_sisvan_counts(
            ["26", "62", "1.02", "1.108"], ["2.35%", "5.6%", "92.06%"], "110001")
        self.assertEqual(values, (26, 62, 1020, 1108))
        self.assertTrue(changed)
        self.assertEqual(solutions, 1)
        values, changed, solutions = analysis.interpret_sisvan_counts(
            ["1", "2", "7", "10"], ["10%", "20%", "70%"], "110001")
        self.assertEqual(values, (1, 2, 7, 10))
        self.assertFalse(changed)
        self.assertEqual(solutions, 1)
        values, changed, solutions = analysis.interpret_sisvan_counts(
            ["0", "0", "6", "6"], ["-", "-", "100%"], "431725")
        self.assertEqual(values, (0, 0, 6, 6))
        self.assertFalse(changed)
        self.assertEqual(solutions, 1)
        for counts in (["1", "2", "7", "11"], ["-1", "2", "7", "8"],
                       ["nan", "2", "7", "9"], ["1.00001", "2", "7", "10"]):
            with self.assertRaises(ValueError):
                analysis.interpret_sisvan_counts(counts, ["10%", "20%", "70%"], "110001")

    def test_coherent_integers_never_enter_the_scale_search(self):
        with patch.object(analysis, "product", side_effect=AssertionError("Não extrapolar linha coerente")):
            result = analysis.interpret_sisvan_counts(
                ["9", "14", "246", "269"], ["3.35%", "5.2%", "91.45%"], "110003")
        self.assertEqual(result, ((9, 14, 246, 269), False, 1))
        self.assertTrue(self.base.sisvan_escalas_compativeis.eq(1).all())


if __name__ == "__main__":
    unittest.main()
