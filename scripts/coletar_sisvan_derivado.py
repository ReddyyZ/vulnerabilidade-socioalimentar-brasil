#!/usr/bin/env python3
"""Coleta agregados municipais do Relatorio Publico do SISVAN.

Recorte padrao: 2025, criancas de 0 a <5 anos, IMC por idade e todos os
demais filtros. Usa a exportacao XLSX oficial por UF, nao microdados.
Dependencias: requests e openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import io
import math
from itertools import product
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

PORTAL = "https://sisaps.saude.gov.br/sisvan/relatoriopublico/"
ENDPOINT = PORTAL + "estadonutricional"
IBGE_UFS = "https://servicodados.ibge.gov.br/api/v1/localidades/estados"
IBGE_MUN = IBGE_UFS + "/{uf}/municipios"
COUNT_COLS = [
    "avaliados", "magreza_acentuada_n", "magreza_n", "desnutricao_n",
    "eutrofia_n", "risco_sobrepeso_n", "sobrepeso_n", "obesidade_n",
]
PCT_COLS = [
    "magreza_acentuada_pct_oficial", "magreza_pct_oficial",
    "desnutricao_pct_oficial", "eutrofia_pct_oficial",
    "risco_sobrepeso_pct_oficial", "sobrepeso_pct_oficial",
    "obesidade_pct_oficial",
]
COLS = [
    "codigo_ibge", "municipio", "uf", "ano", "avaliados",
    "magreza_acentuada_n", "magreza_n", "desnutricao_n", "desnutricao_pct",
    "magreza_acentuada_pct_oficial", "magreza_pct_oficial",
    "desnutricao_pct_oficial", "eutrofia_n", "eutrofia_pct_oficial",
    "risco_sobrepeso_n", "risco_sobrepeso_pct_oficial",
    "sobrepeso_n", "sobrepeso_pct_oficial", "obesidade_n",
    "obesidade_pct_oficial",
]


def new_session(retries):
    policy = Retry(
        total=retries, connect=retries, read=retries, status=retries,
        backoff_factor=1, status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset(("GET", "POST")),
        respect_retry_after_header=True, raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=policy, pool_connections=2, pool_maxsize=2)
    result = requests.Session()
    result.mount("https://", adapter)
    result.mount("http://", adapter)
    result.headers.update({
        "User-Agent": "pesquisa-academica-sisvan/2.0 (+%s)" % PORTAL,
        "Accept-Language": "pt-BR,pt;q=0.9",
    })
    return result


def http(s, method, url, timeout, **kwargs):
    response = s.request(method, url, timeout=timeout, **kwargs)
    if response.status_code == 429:
        raise RuntimeError(
            "HTTP 429; Retry-After=%s" % response.headers.get("Retry-After", "?")
        )
    response.raise_for_status()
    return response


def get_ibge(s, timeout, selected):
    states = http(s, "GET", IBGE_UFS, timeout, params={"orderBy": "id"}).json()
    states.sort(key=lambda x: int(x["id"]))
    if selected:
        states = [x for x in states if x["sigla"].upper() in selected]
    expected = len(selected) if selected else 27
    if len(states) != expected:
        raise RuntimeError("UF inexistente ou resposta IBGE incompleta")
    output = []
    for state in states:
        uf, code = state["sigla"].upper(), str(state["id"])
        data = http(
            s, "GET", IBGE_MUN.format(uf=code), timeout,
            params={"orderBy": "nome"},
        ).json()
        municipalities = []
        for item in data:
            ibge = str(item["id"])
            if not re.fullmatch(r"\d{7}", ibge):
                raise RuntimeError("Codigo IBGE invalido: %r" % ibge)
            municipalities.append({
                "codigo_ibge": ibge, "municipio": item["nome"], "uf": uf,
            })
        output.append((uf, code, municipalities))
    return output


def report_payload(
    year, uf, index_code="4", age_start="0", age_end="5",
):
    # coMunicipioIbge=99 e essencial; vazio devolve somente totais.
    return {
        "excel": "1", "tpRelatorio": "2", "coVisualizacao": "3",
        "nuAno": str(year), "nuMes[]": "99", "tpFiltro": "M",
        "coRegiao": "99", "coUfIbge": uf, "coMunicipioIbge": "99",
        "noRegional": "", "st_cobertura": "99", "nu_ciclo_vida": "1",
        "nu_idade_inicio": str(age_start), "nu_idade_fim": str(age_end),
        "nu_indice_cri": str(index_code),
        "nu_indice_ado": "1", "nu_idade_ges": "99", "ds_sexo2": "1",
        "ds_raca_cor2": "99", "co_sistema_origem": "0",
        "CO_POVO_COMUNIDADE": "TODOS", "CO_ESCOLARIDADE": "TODOS",
        "tpAbrangencia": "M", "tpAbrangenciaEas": "",
    }


class Limiter:
    def __init__(self, interval):
        self.interval, self.last = max(0, interval), 0

    def wait(self):
        remaining = self.interval - (time.monotonic() - self.last)
        if remaining > 0:
            time.sleep(remaining)
        self.last = time.monotonic()


def start_portal(s, timeout, limiter):
    limiter.wait()
    page = http(s, "GET", PORTAL, timeout).text
    expected = [
        'action="/sisvan/relatoriopublico/estadonutricional"',
        'name="coMunicipioIbge"', 'name="nu_indice_cri"',
        'value="3">Altura X Idade', 'value="4">IMC X Idade',
    ]
    missing = [x for x in expected if x not in page]
    if missing:
        raise RuntimeError("Formulario SISVAN mudou; ausentes: " + repr(missing))


def fetch_xlsx(
    s, args, limiter, uf, uf_code, payload=None, debug_name=None,
):
    for attempt in range(args.retries + 1):
        limiter.wait()
        response = http(
            s, "POST", ENDPOINT,
            (args.connect_timeout, args.read_timeout),
            data=payload or report_payload(args.year, uf_code),
            headers={
                "Referer": PORTAL,
                "Accept": "application/vnd.openxmlformats-officedocument."
                          "spreadsheetml.sheet,*/*;q=0.5",
            },
        )
        if response.content.startswith(b"PK\x03\x04"):
            return response.content
        args.debug_dir.mkdir(parents=True, exist_ok=True)
        kind = response.headers.get("Content-Type", "").lower()
        prefix = debug_name or uf
        debug = args.debug_dir / (
            prefix + "_resposta_invalida"
            + (".html" if "html" in kind else ".bin")
        )
        debug.write_bytes(response.content)
        if attempt == args.retries:
            raise RuntimeError("Resposta SISVAN invalida; debug: %s" % debug)
        delay = min(60, 2 ** attempt)
        print("  resposta invalida; retry em %ss; debug: %s" % (delay, debug))
        time.sleep(delay)
    raise AssertionError("tentativas esgotadas")


def parse_count(value):
    """Corrige o XLSX que grava 1.234 pessoas como numero decimal 1.234."""
    if value is None or value == "" or isinstance(value, bool):
        raise ValueError("contagem invalida: %r" % value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value) or value < 0:
            raise ValueError("contagem invalida: %r" % value)
        return int(value) if value.is_integer() else round(value * 1000)
    text = str(value).strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+", text):
        text = text.replace(".", "")
    if not text.isdigit():
        raise ValueError("contagem invalida: %r" % value)
    return int(text)


def count_candidates(value):
    """Possiveis inteiros quando 2.000 foi gravado no XLSX como numero 2."""
    base = parse_count(value)
    integral_source = (
        isinstance(value, int)
        or (isinstance(value, float) and value.is_integer())
        or bool(re.fullmatch(r"\d+", str(value).strip()))
    )
    return [base, base * 1000] if integral_source and base > 0 else [base]


def parse_percent(value):
    text = str(value).strip()
    # O exportador usa "-" quando a categoria tem zero casos.
    if text in {"-", "–", "—"}:
        return 0.0
    number = float(text.replace("%", "").replace(",", "."))
    if not 0 <= number <= 100:
        raise ValueError("percentual invalido: %r" % value)
    return number


def municipal_code(value):
    try:
        number = int(value)
    except (TypeError, ValueError):
        return None
    text = str(number)
    return text if re.fullmatch(r"\d{6}", text) else None


def parse_xlsx(binary, uf):
    book = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    found = {}
    try:
        for source in book.active.iter_rows(values_only=True):
            row = list(source) + [None] * max(0, 18 - len(source))
            code = municipal_code(row[3])
            if code is None or str(row[2]).strip().upper() != uf:
                continue
            if code in found:
                raise RuntimeError("Codigo SISVAN duplicado: %s" % code)
            names = [
                "magreza_acentuada_n", "magreza_n", "eutrofia_n",
                "risco_sobrepeso_n", "sobrepeso_n", "obesidade_n",
            ]
            raw_counts = [row[5], row[7], row[9], row[11], row[13], row[15]]
            pct_names = [
                "magreza_acentuada_pct_oficial", "magreza_pct_oficial",
                "eutrofia_pct_oficial", "risco_sobrepeso_pct_oficial",
                "sobrepeso_pct_oficial", "obesidade_pct_oficial",
            ]
            pct_values = [
                parse_percent(row[6]), parse_percent(row[8]),
                parse_percent(row[10]), parse_percent(row[12]),
                parse_percent(row[14]), parse_percent(row[16]),
            ]
            pcts = dict(zip(pct_names, pct_values))
            total_options = count_candidates(row[17])
            if total_options == [0]:
                if any(parse_count(value) != 0 for value in raw_counts):
                    raise RuntimeError("%s/%s: categorias sem denominador" % (uf, code))
                # Municipio exibido pelo SISVAN, mas sem avaliados no recorte.
                # A juncao com o IBGE o manterá com metricas vazias.
                continue
            solutions = []
            for total_option in total_options:
                for values in product(*(count_candidates(x) for x in raw_counts)):
                    if sum(values) != total_option:
                        continue
                    if all(
                        abs(value / total_option * 100 - pct) <= 0.011
                        for value, pct in zip(values, pct_values)
                    ):
                        solutions.append((total_option, values))
            if not solutions:
                raise RuntimeError("%s/%s: escala das contagens irresolvivel" % (uf, code))
            total, values = min(solutions, key=lambda item: item[0])
            counts = dict(zip(names, values))
            if abs(sum(pcts.values()) - 100) > 0.10:
                raise RuntimeError("%s/%s: percentuais nao somam 100" % (uf, code))
            under = counts["magreza_acentuada_n"] + counts["magreza_n"]
            found[code] = {
                "avaliados": total, **counts, "desnutricao_n": under,
                "desnutricao_pct": round(under / total * 100, 4), **pcts,
                "desnutricao_pct_oficial": round(
                    pcts["magreza_acentuada_pct_oficial"]
                    + pcts["magreza_pct_oficial"], 2,
                ),
            }
    finally:
        book.close()
    if not found:
        raise RuntimeError("Nenhum municipio SISVAN reconhecido para " + uf)
    return found


def join_ibge(source, municipalities, year, uf):
    mapping = {x["codigo_ibge"][:6]: x for x in municipalities}
    if len(mapping) != len(municipalities):
        raise RuntimeError("Colisao em codigos IBGE sem digito em " + uf)
    unknown = sorted(set(source) - set(mapping))
    if unknown:
        raise RuntimeError("Codigos SISVAN sem IBGE em %s: %s" % (uf, unknown))
    blanks = {
        x: "" for x in COUNT_COLS + ["desnutricao_pct"] + PCT_COLS
    }
    rows = []
    for municipality in municipalities:
        ibge = municipality["codigo_ibge"]
        row = {
            "codigo_ibge": ibge, "municipio": municipality["municipio"],
            "uf": uf, "ano": year,
        }
        row.update(source.get(ibge[:6], blanks))
        rows.append(row)
    return rows


def validate(rows):
    seen, missing, max_difference = set(), [], 0
    for row in rows:
        code = str(row["codigo_ibge"])
        if not re.fullmatch(r"\d{7}", code) or code in seen:
            raise RuntimeError("Codigo IBGE invalido/duplicado: " + code)
        seen.add(code)
        if str(row["avaliados"]).strip() == "":
            if any(str(row[x]).strip() for x in COLS[5:]):
                raise RuntimeError("Dados parciais em municipio ausente: " + code)
            missing.append(row)
            continue
        total = int(row["avaliados"])
        cats = [
            int(row["magreza_acentuada_n"]), int(row["magreza_n"]),
            int(row["eutrofia_n"]), int(row["risco_sobrepeso_n"]),
            int(row["sobrepeso_n"]), int(row["obesidade_n"]),
        ]
        if total <= 0 or total < sum(cats):
            raise RuntimeError("avaliados < categorias em " + code)
        if int(row["desnutricao_n"]) != cats[0] + cats[1]:
            raise RuntimeError("desnutricao_n incorreta em " + code)
        calculated = float(row["desnutricao_pct"])
        if abs(calculated - (cats[0] + cats[1]) / total * 100) > 0.00011:
            raise RuntimeError("desnutricao_pct incorreta em " + code)
        for field in PCT_COLS:
            if not 0 <= float(row[field]) <= 100:
                raise RuntimeError("percentual fora de 0..100 em " + code)
        max_difference = max(
            max_difference,
            abs(calculated - float(row["desnutricao_pct_oficial"])),
        )
    return {
        "total": len(rows), "with_data": len(rows) - len(missing),
        "missing": missing, "max_difference": max_difference,
    }


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLS)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def load_checkpoint(path, year, uf, municipalities):
    if not path.exists():
        return None
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames != COLS:
                return None
            rows = list(reader)
        expected = {x["codigo_ibge"] for x in municipalities}
        if {x["codigo_ibge"] for x in rows} != expected:
            return None
        if any(x["ano"] != str(year) or x["uf"] != uf for x in rows):
            return None
        validate(rows)
        return rows
    except (OSError, ValueError, RuntimeError):
        return None


def write_readme(args, stats):
    missing = stats["missing"]
    missing_text = ""
    if missing:
        listed = "\n".join(
            "- %s — %s/%s" % (x["codigo_ibge"], x["municipio"], x["uf"])
            for x in missing[:100]
        )
        missing_text = "\n### Municipios sem dados\n\n" + listed
        if len(missing) > 100:
            missing_text += "\n- ..."
    text = f"""# SISVAN municipal — estado nutricional ({args.year})

Arquivo: sisvan_municipios_{args.year}.csv

## Fonte e metodo

- Fonte: Relatorio Publico oficial do SISVAN: {PORTAL}
- Endpoint: POST {ENDPOINT}
- Resposta: exportacao oficial XLSX, uma requisicao por UF.
- Municipios e chave de sete digitos: API de Localidades do IBGE: {IBGE_UFS}
- Coleta: {datetime.now().astimezone().isoformat(timespec="seconds")}

Foi usada a exportacao agregada, nao microdados. A API de Dados Abertos e
paginada em registros e seu catalogo descreve bases individualizadas que podem
divergir do Relatorio Publico, pois este consolida o ultimo acompanhamento.

A sessao HTTP conserva PHPSESSID e cookies do balanceador. Nao havia token CSRF.
O reCAPTCHA de estado nutricional estava comentado; nenhum CAPTCHA foi burlado.

## Recorte e parametros

- Ano {args.year}; nuMes[]=99 (ano inteiro).
- Municipio: tpFiltro=M, coMunicipioIbge=99 (todos), uma UF por vez.
- Criancas de 0 a <5 anos: nu_ciclo_vida=1, idades 0 e 5.
- IMC por idade: nu_indice_cri=4.
- Sexo, raca/cor, origem, povo/comunidade e escolaridade: todos.

O codigo SISVAN de seis digitos foi ligado aos seis primeiros digitos do codigo
IBGE. O CSV preserva a chave oficial IBGE de sete digitos.

## Indicador

desnutricao_n = magreza_acentuada_n + magreza_n

desnutricao_pct = desnutricao_n / avaliados * 100

desnutricao_pct_oficial soma os percentuais oficiais das duas categorias, que
sao arredondados separadamente. As demais colunas _pct_oficial preservam os
valores exibidos pelo SISVAN.

## Resultado e verificacoes

- Municipios IBGE: **{stats["total"]}**.
- Com dados SISVAN: **{stats["with_data"]}**.
- Sem dados SISVAN: **{len(missing)}**.
- Diferenca ante a referencia solicitada de 5.570: **{stats["total"] - 5570:+d}**.
- Maior diferenca calculado versus oficial: **{stats["max_difference"]:.4f} p.p.**

A API atual do IBGE define o universo, que pode refletir municipio criado
recentemente. Ausencias ficam vazias, nunca zero. Foram validados codigo IBGE,
duplicatas, soma das categorias, formulas e limites dos percentuais.
{missing_text}

## Reproducao e retomada

    python3 scripts/coletar_sisvan.py --year {args.year}

Checkpoints ficam em data/sisvan/checkpoints/{args.year}/UF.csv. Uma nova
execucao retoma os validos. Use --force para recarregar. O coletor e sequencial,
tem intervalo conservador, timeouts, Retry-After e retries exponenciais.
"""
    args.readme.parent.mkdir(parents=True, exist_ok=True)
    args.readme.write_text(text, encoding="utf-8")


def collect(args):
    timeout = (args.connect_timeout, args.read_timeout)
    states = get_ibge(new_session(args.retries), timeout, args.selected_ufs)
    expected_total = sum(len(x[2]) for x in states)
    print("Municipios IBGE no universo selecionado:", expected_total)
    sisvan = new_session(args.retries)
    limiter = Limiter(args.delay)
    started = False
    rows_all = []
    for index, (uf, uf_code, municipalities) in enumerate(states, 1):
        print("[%02d/%02d] %s" % (index, len(states), uf))
        path = args.checkpoint_dir / str(args.year) / (uf + ".csv")
        rows = None if args.force else load_checkpoint(
            path, args.year, uf, municipalities
        )
        if rows is not None:
            print("  checkpoint valido reutilizado")
        else:
            if not started:
                start_portal(sisvan, timeout, limiter)
                started = True
            print("  baixando exportacao XLSX agregada...")
            binary = fetch_xlsx(sisvan, args, limiter, uf, uf_code)
            rows = join_ibge(parse_xlsx(binary, uf), municipalities, args.year, uf)
            validate(rows)
            write_csv(path, rows)
            print("  checkpoint salvo:", path)
        stats = validate(rows)
        print("  municipios encontrados:", len(municipalities))
        print("  com dados SISVAN:", stats["with_data"])
        print("  sem dados:", len(stats["missing"]))
        rows_all.extend(rows)
    rows_all.sort(key=lambda x: int(x["codigo_ibge"]))
    stats = validate(rows_all)
    if stats["total"] != expected_total:
        raise RuntimeError("Total final diverge do universo IBGE")
    write_csv(args.output, rows_all)
    write_readme(args, stats)
    print("\nTOTAL")
    print("municipios IBGE:", stats["total"])
    print("municipios com dados SISVAN:", stats["with_data"])
    print("municipios sem dados:", len(stats["missing"]))
    print("CSV:", args.output)
    print("Relatorio:", args.readme)


def main():
    parser = argparse.ArgumentParser(
        description="Coleta agregados municipais do Relatorio Publico SISVAN."
    )
    parser.add_argument("--year", type=int, default=2025)
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--checkpoint-dir", type=Path, default=Path("data/sisvan/checkpoints")
    )
    parser.add_argument("--debug-dir", type=Path, default=Path("debug/sisvan"))
    parser.add_argument("--readme", type=Path, default=Path("data/sisvan/README.md"))
    parser.add_argument("--delay", type=float, default=1.0)
    parser.add_argument("--connect-timeout", type=float, default=20.0)
    parser.add_argument("--read-timeout", type=float, default=180.0)
    parser.add_argument("--retries", type=int, default=4)
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--ufs", help="Siglas separadas por virgula, ex.: AC,SP")
    args = parser.parse_args()
    if args.delay < 0 or args.connect_timeout <= 0 or args.read_timeout <= 0:
        parser.error("delay >= 0 e timeouts positivos")
    if args.retries < 0:
        parser.error("retries >= 0")
    args.selected_ufs = (
        frozenset(x.strip().upper() for x in args.ufs.split(",") if x.strip())
        if args.ufs else None
    )
    if args.selected_ufs and any(
        not re.fullmatch(r"[A-Z]{2}", x) for x in args.selected_ufs
    ):
        parser.error("Use siglas como AC,SP")
    args.output = args.output or Path(
        "data/sisvan/sisvan_municipios_%s.csv" % args.year
    )
    collect(args)


if __name__ == "__main__":
    main()
