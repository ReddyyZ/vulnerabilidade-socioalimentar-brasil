#!/usr/bin/env python3
"""Coleta altura e peso por idade de crianças de 0 a menos de cinco anos."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

if __package__:
    from . import sisvan_cliente as helper
else:
    import sisvan_cliente as helper


UF_CODES = [
    ("RO", "11"), ("AC", "12"), ("AM", "13"), ("RR", "14"),
    ("PA", "15"), ("AP", "16"), ("TO", "17"), ("MA", "21"),
    ("PI", "22"), ("CE", "23"), ("RN", "24"), ("PB", "25"),
    ("PE", "26"), ("AL", "27"), ("SE", "28"), ("BA", "29"), ("MG", "31"),
    ("ES", "32"), ("RJ", "33"), ("SP", "35"), ("PR", "41"),
    ("SC", "42"), ("RS", "43"), ("MS", "50"), ("MT", "51"),
    ("GO", "52"), ("DF", "53"),
]
BASE_COLUMNS = ("Região", "Código UF", "UF", "Código IBGE", "Município")
AGE_KEY = "0_a_menor_5_anos"
AGE_LABEL = "0 a < 5 anos"
COLLECTOR_VERSION = "6.0"
MANIFEST_COLUMNS = (
    "id_coleta", "fonte", "arquivo_local", "nome_original",
    "indice_antropometrico", "faixa_etaria", "ano_inicio", "ano_fim",
    "abrangencia", "filtros", "data_hora_coleta", "quantidade_registros",
    "hash_arquivo", "status", "versao_coletor", "observacoes",
)


@dataclass(frozen=True)
class Indicator:
    key: str
    code: str
    official_title: str
    categories: tuple[str, ...]


@dataclass(frozen=True)
class Query:
    indicator: Indicator
    year: int


@dataclass(frozen=True)
class Schema:
    categories: tuple[str, ...]
    columns: tuple[str, ...]

    @property
    def count_columns(self):
        return tuple(category + " - Quantidade" for category in self.categories)

    @property
    def percent_columns(self):
        return tuple(category + " - %" for category in self.categories)


INDICATORS = {
    "altura_por_idade": Indicator(
        "altura_por_idade", "3", "ALTURA X IDADE",
        ("Altura Muito Baixa para a Idade", "Altura Baixa para a Idade",
         "Altura Adequada para a Idade"),
    ),
    "peso_por_idade": Indicator(
        "peso_por_idade", "1", "PESO X IDADE",
        ("Peso Muito Baixo para a Idade", "Peso Baixo para a Idade",
         "Peso Adequado ou Eutrófico", "Peso Elevado para a Idade"),
    ),
}


def normalize_header(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def read_schema(sheet, indicator):
    headers = [[normalize_header(value) for value in row]
               for row in sheet.iter_rows(min_row=1, max_row=15, values_only=True)]
    if not any(indicator.official_title in row for row in headers):
        raise RuntimeError("Resposta não corresponde a " + indicator.official_title)
    header = next((row for row in headers
                   if all(name in row for name in BASE_COLUMNS + ("Total",))), None)
    if header is None:
        raise RuntimeError("Cabeçalho oficial do SISVAN não reconhecido")
    positions = [header.index(name) for name in BASE_COLUMNS]
    if positions != list(range(positions[0], positions[0] + 5)):
        raise RuntimeError("Colunas territoriais mudaram no XLSX do SISVAN")
    total_position = header.index("Total", positions[-1] + 1)
    categories = tuple(value for value in header[positions[-1] + 1:total_position] if value)
    if categories != indicator.categories:
        raise RuntimeError(f"Categorias oficiais inesperadas para {indicator.key}: {categories}")
    columns = list(BASE_COLUMNS)
    for category in categories:
        columns.extend((category + " - Quantidade", category + " - %"))
    return Schema(categories, tuple(columns) + ("Total",))


def parse_export(binary, expected_uf, indicator):
    book = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    rows, seen = [], set()
    try:
        schema = read_schema(book.active, indicator)
        width = len(schema.columns)
        for source in book.active.iter_rows(values_only=True):
            values = list(source) + [None] * max(0, width - len(source))
            code = helper.municipal_code(values[3])
            uf = normalize_header(values[2])
            if code is None or uf != expected_uf:
                continue
            if code in seen:
                raise RuntimeError(f"Código SISVAN duplicado em {uf}: {code}")
            seen.add(code)
            row = dict(zip(BASE_COLUMNS, (
                normalize_header(values[0]), str(int(values[1])).zfill(2),
                uf, code, normalize_header(values[4]),
            )))
            for index, category in enumerate(schema.categories):
                row[category + " - Quantidade"] = values[5 + 2 * index]
                percentage = str(values[6 + 2 * index]).strip()
                if percentage not in {"-", "–", "—"} and not percentage.endswith("%"):
                    raise ValueError("Percentual SISVAN inesperado: " + percentage)
                row[category + " - %"] = percentage
            row["Total"] = values[width - 1]
            rows.append(row)
    finally:
        book.close()
    if not rows:
        raise RuntimeError("Nenhuma linha municipal encontrada para " + expected_uf)
    return rows, schema


def validate(rows, schema):
    seen, zeros = set(), []
    for row in rows:
        code = str(row["Código IBGE"])
        if not re.fullmatch(r"\d{6}", code) or code in seen:
            raise RuntimeError("Código SISVAN inválido ou duplicado: " + code)
        seen.add(code)
        counts = [float(row[column]) for column in schema.count_columns]
        percentages = [helper.parse_percent(row[column]) for column in schema.percent_columns]
        total = float(row["Total"])
        if any(not math.isfinite(value) or value < 0 for value in counts + [total]):
            raise RuntimeError("Contagem negativa ou não finita em " + code)
        if total == 0:
            if any(counts) or any(percentages):
                raise RuntimeError("Linha zero inconsistente em " + code)
            zeros.append(row)
        elif abs(sum(percentages) - 100) > 0.10:
            raise RuntimeError("Percentuais não somam 100 em " + code)
    # Contagens decimais do exportador são preservadas, não corrigidas aqui.
    return {"total": len(rows), "zeros": zeros}


def atomic_write_bytes(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(content)
    temporary.replace(path)


def write_csv(path, rows, columns, lineterminator="\r\n"):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, lineterminator=lineterminator)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def sha256(content):
    return hashlib.sha256(content).hexdigest()


def portable_path(path):
    try:
        return str(path.resolve().relative_to(Path.cwd().resolve()))
    except ValueError:
        return str(path.resolve())


class Manifest:
    def __init__(self, path):
        self.path, self.rows = path, {}
        if path.exists():
            with path.open(newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                if tuple(reader.fieldnames or ()) != MANIFEST_COLUMNS:
                    raise RuntimeError("Esquema inesperado no manifesto: " + str(path))
                for row in reader:
                    if row["id_coleta"] in self.rows:
                        raise RuntimeError("Identificador duplicado no manifesto: " + row["id_coleta"])
                    self.rows[row["id_coleta"]] = row

    def record(self, row, reused=False):
        previous = self.rows.get(row["id_coleta"])
        if reused and previous:
            if previous["hash_arquivo"] != row["hash_arquivo"]:
                raise RuntimeError("Arquivo bruto diverge do manifesto: " + row["arquivo_local"])
            row["data_hora_coleta"] = previous["data_hora_coleta"]
        self.rows[row["id_coleta"]] = row
        ordered = sorted(self.rows.values(), key=lambda item: (
            item["ano_inicio"], item["indice_antropometrico"],
            item["faixa_etaria"], item["arquivo_local"],
        ))
        write_csv(self.path, ordered, MANIFEST_COLUMNS, lineterminator="\n")


def parse_list(value):
    return list(dict.fromkeys(item.strip() for item in value.split(",") if item.strip()))


def resolve_queries(args):
    indices = parse_list(args.indices)
    if not indices or any(index not in INDICATORS for index in indices):
        raise RuntimeError("Índices permitidos: altura_por_idade,peso_por_idade")
    year = args.year
    if type(year) is not int or year <= 0:
        raise RuntimeError("O ano deve ser um inteiro positivo")
    return [Query(INDICATORS[index], year) for index in indices]


def selected_states(value):
    if value is None:
        return UF_CODES
    selected = set(parse_list(value.upper()))
    invalid = selected - {uf for uf, _ in UF_CODES}
    if not selected or invalid:
        raise RuntimeError("UF inválida: " + ", ".join(sorted(invalid or {value})))
    return [state for state in UF_CODES if state[0] in selected]


def raw_path(args, query, uf):
    return args.raw_dir / query.indicator.key / AGE_KEY / str(query.year) / "ufs" / (uf + ".xlsx")


def output_path(args, query, states):
    suffix = "" if len(states) == len(UF_CODES) else "_ufs_" + "-".join(uf for uf, _ in states)
    name = f"sisvan_municipios_{query.indicator.key}_{AGE_KEY}_{query.year}{suffix}.csv"
    return args.output or args.output_dir / query.indicator.key / name


def write_metadata(output, query, stats, states, schema, originals):
    data = {
        "fonte": helper.PORTAL, "endpoint": helper.ENDPOINT,
        "arquivo": portable_path(output),
        "fase": "crianca", "fase_oficial": "CRIANÇA", "codigo_fase_sisvan": "1",
        "indice": query.indicator.key, "indice_oficial": query.indicator.official_title,
        "codigo_indice_sisvan": query.indicator.code,
        "faixa_etaria": AGE_KEY, "faixa_etaria_descricao": AGE_LABEL,
        "nu_idade_inicio": "0", "nu_idade_fim": "5",
        "ano": query.year, "ufs": [uf for uf, _ in states],
        "colunas": list(schema.columns), "linhas": stats["total"],
        "linhas_total_zero": len(stats["zeros"]), "versao_coletor": COLLECTOR_VERSION,
        "sha256_csv": sha256(output.read_bytes()), "xlsx_originais": originals,
        "contagens": "valores das celulas preservados; sem normalizacao de escala",
        "indicadores_derivados": [],
        "gerado_em": datetime.now().astimezone().isoformat(timespec="seconds"),
        "observacao": (
            "CSV convertido dos XLSX oficiais por UF. Cabecalhos multinivel "
            "sao achatados em categoria - Quantidade e categoria - %. "
            "Valores das celulas nao sao corrigidos ou recalculados. "
            "Os XLSX preservados sao os arquivos originais exatos."
        ),
    }
    atomic_write_bytes(output.with_suffix(".metadados.json"),
                       (json.dumps(data, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))


def collect_query(args, query, states, session, limiter, manifest):
    print(f"\nCONSULTA: CRIANÇA | {query.indicator.official_title} | {AGE_LABEL} | {query.year}")
    all_rows, expected_schema, originals = [], None, []
    for position, (uf, uf_code) in enumerate(states, 1):
        print(f"[{position:02d}/{len(states):02d}] {uf}")
        raw = raw_path(args, query, uf)
        payload = helper.report_payload(query.year, uf_code, query.indicator.code)
        reused = raw.exists() and not args.force
        if reused:
            content = raw.read_bytes()
        else:
            content = helper.fetch_xlsx(session, args, limiter, payload,
                                        f"{query.indicator.key}_{query.year}_{uf}")
            atomic_write_bytes(raw, content)
        rows, schema = parse_export(content, uf, query.indicator)
        validate(rows, schema)
        if expected_schema is not None and schema != expected_schema:
            raise RuntimeError("Esquema divergente entre UFs: " + query.indicator.key)
        expected_schema = schema
        all_rows.extend(rows)
        entry = {
            "id_coleta": f"sisvan-{query.year}-{query.indicator.key}-{AGE_KEY}-{uf}",
            "fonte": helper.PORTAL, "arquivo_local": portable_path(raw), "nome_original": "",
            "indice_antropometrico": query.indicator.official_title, "faixa_etaria": AGE_LABEL,
            "ano_inicio": str(query.year), "ano_fim": str(query.year),
            "abrangencia": "Municipios de " + uf,
            "filtros": json.dumps(payload, ensure_ascii=False, sort_keys=True),
            "data_hora_coleta": datetime.now().astimezone().isoformat(timespec="seconds"),
            "quantidade_registros": str(len(rows)), "hash_arquivo": sha256(content),
            "status": "validado", "versao_coletor": COLLECTOR_VERSION,
            "observacoes": "XLSX oficial preservado sem alteracao",
        }
        manifest.record(entry, reused=reused)
        originals.append({key: entry[key] for key in (
            "arquivo_local", "hash_arquivo", "data_hora_coleta", "filtros", "quantidade_registros",
        )})
        print(f"  XLSX {'reutilizado' if reused else 'salvo'}: {raw}; municípios: {len(rows)}")
    all_rows.sort(key=lambda row: int(row["Código IBGE"]))
    stats = validate(all_rows, expected_schema)
    output = output_path(args, query, states)
    write_csv(output, all_rows, expected_schema.columns)
    write_metadata(output, query, stats, states, expected_schema, originals)
    print(f"  CSV: {output}; linhas: {stats['total']}; Total=0: {len(stats['zeros'])}")
    return output


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, default=2025, help="Ano; padrão: 2025")
    parser.add_argument("--indices", default=",".join(INDICATORS),
                        help="altura_por_idade,peso_por_idade; padrão: ambos")
    parser.add_argument("--ufs", help="UFs separadas por vírgulas; padrão: todas as 27")
    parser.add_argument("--raw-dir", type=Path, default=Path("dados/coletas/sisvan/brutos"))
    parser.add_argument("--output-dir", type=Path, default=Path("dados/coletas/sisvan/convertidos"))
    parser.add_argument("--output", type=Path, help="CSV de destino; exige selecionar um único índice")
    parser.add_argument("--manifest", type=Path, default=Path("dados/coletas/sisvan/manifesto_coletas.csv"))
    parser.add_argument("--debug-dir", type=Path, default=Path("dados/coletas/sisvan/debug"))
    parser.add_argument("--delay", type=float, default=1.0)
    parser.add_argument("--connect-timeout", type=float, default=20.0)
    parser.add_argument("--read-timeout", type=float, default=180.0)
    parser.add_argument("--retries", type=int, default=4)
    parser.add_argument("--force", action="store_true", help="Baixa novamente mesmo quando o XLSX já existe")
    parser.add_argument("--dry-run", action="store_true", help="Mostra as consultas sem acessar a rede ou gravar")
    return parser


def main():
    parser = build_parser()
    args = parser.parse_args()
    if (not all(math.isfinite(value) for value in (args.delay, args.connect_timeout, args.read_timeout))
            or args.delay < 0 or args.connect_timeout <= 0 or args.read_timeout <= 0 or args.retries < 0):
        parser.error("Delay e retries devem ser não negativos; timeouts devem ser positivos e finitos")
    try:
        queries = resolve_queries(args)
        states = selected_states(args.ufs)
        if args.output and len(queries) != 1:
            raise RuntimeError("--output exige selecionar um único índice")
        root = Path(__file__).resolve().parents[1]
        protected = [root / folder for folder in ("dados/pesquisa", "dados/brutos", "metadados")]
        for path in (args.raw_dir, args.output_dir, args.manifest, args.debug_dir, args.output):
            if path is not None and (any(path.resolve().is_relative_to(p.resolve()) for p in protected)
                                     or path.resolve() == root / "experimento.ipynb"):
                raise RuntimeError("Escolha destinos fora das entradas, origens e notebook preservados")
    except (RuntimeError, ValueError) as exc:
        parser.error(str(exc))
    print("Consultas planejadas:")
    for query in queries:
        print(f"- CRIANÇA | {query.indicator.official_title} | {AGE_LABEL} | {query.year}")
    print("UFs:", ", ".join(uf for uf, _ in states))
    if args.dry_run:
        return
    manifest = Manifest(args.manifest)
    session = helper.new_session(args.retries)
    try:
        limiter = helper.Limiter(args.delay)
        if any(args.force or not raw_path(args, query, uf).exists()
               for query in queries for uf, _ in states):
            helper.start_portal(session, (args.connect_timeout, args.read_timeout), limiter)
        for query in queries:
            collect_query(args, query, states, session, limiter, manifest)
    finally:
        session.close()
    print("\nCOLETA CONCLUÍDA. Manifesto:", args.manifest)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nColeta interrompida; arquivos já gravados foram preservados.", file=sys.stderr)
        raise SystemExit(130)
